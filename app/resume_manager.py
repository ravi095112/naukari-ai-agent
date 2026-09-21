from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import re
import sys

from docx import Document


BASE_DIR = Path(__file__).resolve().parent.parent

RESUME_FILE = BASE_DIR / "Ravi_Kumar_Updated_Resume.docx"

RESUME_DATA_DIR = BASE_DIR / "data" / "resume"
BASELINE_FILE = RESUME_DATA_DIR / "master_resume_baseline.json"
ANALYSIS_FILE = RESUME_DATA_DIR / "resume_change_analysis.json"


MEANINGFUL_KEYWORDS = {
    "experience": [
        "experience",
        "present",
        "joined",
        "promoted",
        "role",
        "designation",
        "responsibil",
        "worked",
        "developed",
        "built",
        "implemented",
        "designed",
        "architect",
    ],

    "technology": [
        "java",
        "spring",
        "spring boot",
        "spring ai",
        "langchain",
        "python",
        "fastapi",
        "microservices",
        "rest api",
        "docker",
        "kubernetes",
        "kafka",
        "aws",
        "azure",
        "rag",
        "retrieval augmented generation",
        "generative ai",
        "llm",
        "machine learning",
    ],

    "achievement": [
        "reduced",
        "increased",
        "improved",
        "optimized",
        "saved",
        "automated",
        "performance",
        "productivity",
        "%",
    ],

    "education": [
        "education",
        "degree",
        "mca",
        "bca",
        "b.tech",
        "btech",
        "master",
        "bachelor",
        "university",
        "college",
    ],

    "certification": [
        "certification",
        "certified",
        "certificate",
        "course",
        "credential",
    ],

    "project": [
        "project",
        "application",
        "platform",
        "system",
        "chatbot",
        "assistant",
    ],
}


def normalize_text(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


def calculate_fingerprint(text: str) -> str:
    return hashlib.sha256(
        text.encode("utf-8")
    ).hexdigest()


def read_resume() -> dict:

    if not RESUME_FILE.exists():
        raise FileNotFoundError(
            f"Master resume not found:\n{RESUME_FILE}"
        )

    document = Document(str(RESUME_FILE))

    paragraphs = [
        normalize_text(p.text)
        for p in document.paragraphs
        if normalize_text(p.text)
    ]

    tables = []

    for table in document.tables:

        rows = []

        for row in table.rows:

            cells = [
                normalize_text(cell.text)
                for cell in row.cells
            ]

            if any(cells):
                rows.append(cells)

        if rows:
            tables.append(rows)

    full_text_parts = paragraphs.copy()

    for table in tables:
        for row in table:
            full_text_parts.append(
                " | ".join(row)
            )

    full_text = normalize_text(
        "\n".join(full_text_parts)
    )

    return {
        "file_name": RESUME_FILE.name,
        "file_path": str(RESUME_FILE),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "paragraphs": paragraphs,
        "tables": tables,
        "full_text": full_text,
    }


def load_baseline() -> dict:

    if not BASELINE_FILE.exists():
        raise FileNotFoundError(
            f"Resume baseline not found:\n"
            f"{BASELINE_FILE}\n\n"
            "Run Step 1 first."
        )

    with BASELINE_FILE.open(
            "r",
            encoding="utf-8",
    ) as f:
        return json.load(f)


def save_baseline(resume: dict):

    fingerprint = calculate_fingerprint(
        resume["full_text"]
    )

    baseline = {
        "created_at": datetime.now(
            timezone.utc
        ).isoformat(),

        "source_file": resume["file_name"],

        "fingerprint": fingerprint,

        "paragraph_count":
            resume["paragraph_count"],

        "table_count":
            resume["table_count"],

        "full_text":
            resume["full_text"],
    }

    RESUME_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with BASELINE_FILE.open(
            "w",
            encoding="utf-8",
    ) as f:

        json.dump(
            baseline,
            f,
            indent=2,
            ensure_ascii=False,
        )


def detect_changes(
        old_text: str,
        new_text: str,
) -> list:

    old_lines = old_text.splitlines()
    new_lines = new_text.splitlines()

    changes = []

    max_lines = max(
        len(old_lines),
        len(new_lines),
    )

    for index in range(max_lines):

        old_line = (
            old_lines[index]
            if index < len(old_lines)
            else ""
        )

        new_line = (
            new_lines[index]
            if index < len(new_lines)
            else ""
        )

        if old_line != new_line:

            changes.append({
                "line": index + 1,
                "old": old_line,
                "new": new_line,
            })

    return changes

def keyword_matches(keyword: str, text: str) -> bool:

    if " " in keyword:
        return keyword in text

    return re.search(
        rf"\b{re.escape(keyword)}\b",
        text,
    ) is not None

def classify_change(change: dict) -> dict:

    old_text = change["old"].lower()
    new_text = change["new"].lower()

    combined_text = (
            old_text + " " + new_text
    )

    categories = []

    for category, keywords in MEANINGFUL_KEYWORDS.items():

        matched_keywords = [
            keyword
            for keyword in keywords
            if keyword_matches(keyword, combined_text)
        ]

        if matched_keywords:

            categories.append({
                "category": category,
                "keywords": matched_keywords,
            })

    old_compact = re.sub(
        r"\s+",
        "",
        old_text,
    )

    new_compact = re.sub(
        r"\s+",
        "",
        new_text,
    )

    if (
            old_compact == new_compact
            and old_text != new_text
    ):

        return {
            "classification": "FORMAT_ONLY",
            "categories": [],
            "reason": (
                "Only whitespace or formatting-related "
                "text differences detected."
            ),
        }

    if categories:

        return {
            "classification": "MEANINGFUL",
            "categories": categories,
            "reason": (
                "Changed content contains career, "
                "technology, project, achievement, "
                "education, or certification information."
            ),
        }

    return {
        "classification": "REVIEW_REQUIRED",
        "categories": [],
        "reason": (
            "Resume content changed, but the change "
            "could not be confidently classified."
        ),
    }


def analyze_changes(changes: list) -> list:

    results = []

    for change in changes:

        classification = classify_change(
            change
        )

        results.append({
            **change,
            **classification,
        })

    return results


def save_analysis(
        current_fingerprint: str,
        baseline_fingerprint: str,
        changes: list,
        analysis: list,
):

    meaningful_count = sum(
        1
        for item in analysis
        if item["classification"]
        == "MEANINGFUL"
    )

    review_count = sum(
        1
        for item in analysis
        if item["classification"]
        == "REVIEW_REQUIRED"
    )

    format_only_count = sum(
        1
        for item in analysis
        if item["classification"]
        == "FORMAT_ONLY"
    )

    result = {
        "analyzed_at": datetime.now(
            timezone.utc
        ).isoformat(),

        "resume_file": RESUME_FILE.name,

        "baseline_fingerprint":
            baseline_fingerprint,

        "current_fingerprint":
            current_fingerprint,

        "changed": bool(changes),

        "change_count":
            len(changes),

        "meaningful_change_count":
            meaningful_count,

        "review_required_count":
            review_count,

        "format_only_count":
            format_only_count,

        "approval_status":
            "PENDING",

        "naukri_upload_approved":
            False,

        "changes":
            analysis,
    }

    RESUME_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with ANALYSIS_FILE.open(
            "w",
            encoding="utf-8",
    ) as f:

        json.dump(
            result,
            f,
            indent=2,
            ensure_ascii=False,
        )

    return result


def load_analysis() -> dict:

    if not ANALYSIS_FILE.exists():

        raise FileNotFoundError(
            f"Analysis file not found:\n"
            f"{ANALYSIS_FILE}\n\n"
            "Run analysis first."
        )

    with ANALYSIS_FILE.open(
            "r",
            encoding="utf-8",
    ) as f:

        return json.load(f)


def approve_current_change():

    analysis = load_analysis()

    # --------------------------------------------------------
    # Validate pending analysis
    # --------------------------------------------------------

    if not analysis.get("changed"):

        print(
            "\nNo resume change exists to approve."
        )

        return

    meaningful_count = analysis.get(
        "meaningful_change_count",
        0,
    )

    review_count = analysis.get(
        "review_required_count",
        0,
    )

    format_only_count = analysis.get(
        "format_only_count",
        0,
    )

    if (
            meaningful_count == 0
            and review_count == 0
    ):

        print(
            "\nNo substantive resume change "
            "is available for approval."
        )

        return

    # --------------------------------------------------------
    # Display pending changes
    # --------------------------------------------------------

    print("\n========================================")
    print("PENDING RESUME CHANGES")
    print("========================================")

    print(
        f"\nMeaningful changes : {meaningful_count}"
    )

    print(
        f"Review required    : {review_count}"
    )

    print(
        f"Format-only changes: {format_only_count}"
    )

    print("\n----------------------------------------")

    changes = analysis.get(
        "changes",
        [],
    )

    for index, item in enumerate(
            changes,
            start=1,
    ):

        print(
            f"\nCHANGE #{index}"
        )

        print(
            f"Line: "
            f"{item.get('line', 'N/A')}"
        )

        print(
            f"Classification: "
            f"{item.get('classification', 'UNKNOWN')}"
        )

        print(
            f"Reason: "
            f"{item.get('reason', '')}"
        )

        print("\nOLD:")

        if item.get("old"):
            print(
                item["old"]
            )
        else:
            print(
                "<new line>"
            )

        print("\nNEW:")

        if item.get("new"):
            print(
                item["new"]
            )
        else:
            print(
                "<removed line>"
            )

        categories = item.get(
            "categories",
            [],
        )

        if categories:

            print("\nCategories:")

            for category in categories:

                print(
                    f"  - "
                    f"{category.get('category')}: "
                    f"{', '.join(category.get('keywords', []))}"
                )

        print(
            "\n----------------------------------------"
        )

    # --------------------------------------------------------
    # Explicit human confirmation
    # --------------------------------------------------------

    print("\n========================================")
    print("NAUKRI UPLOAD APPROVAL")
    print("========================================")

    print(
        "\nReview the changes above carefully."
    )

    print(
        "\nThis approval will allow the future "
        "Naukri upload module to use this resume."
    )

    print(
        "\nType exactly YES to approve."
    )

    print(
        "Any other input will cancel the operation."
    )

    confirmation = input(
        "\nDo you approve these changes "
        "for Naukri upload? Type YES to continue: "
    ).strip()

    if confirmation != "YES":

        print("\n========================================")
        print("APPROVAL CANCELLED")
        print("========================================")

        print(
            "\nNo changes were approved."
        )

        print(
            "\nnaukri_upload_approved remains False."
        )

        print(
            "\nResume baseline was NOT updated."
        )

        return

    # --------------------------------------------------------
    # Approval confirmed
    # --------------------------------------------------------

    analysis[
        "approval_status"
    ] = "APPROVED"

    analysis[
        "naukri_upload_approved"
    ] = True

    analysis[
        "approved_at"
    ] = datetime.now(
        timezone.utc
    ).isoformat()

    analysis[
        "approval_confirmation"
    ] = "YES"

    with ANALYSIS_FILE.open(
            "w",
            encoding="utf-8",
    ) as f:

        json.dump(
            analysis,
            f,
            indent=2,
            ensure_ascii=False,
        )

    # --------------------------------------------------------
    # Update baseline only after approval
    # --------------------------------------------------------

    resume = read_resume()

    save_baseline(resume)

    # --------------------------------------------------------
    # Final status
    # --------------------------------------------------------

    print("\n========================================")
    print("RESUME CHANGE APPROVED")
    print("========================================")

    print(
        "\nApproval recorded successfully."
    )

    print(
        "\nApproved at:"
    )

    print(
        analysis["approved_at"]
    )

    print(
        "\nnaukri_upload_approved: True"
    )

    print(
        "\nBaseline updated to the approved resume."
    )

    print(
        "\nIMPORTANT:"
    )

    print(
        "Naukri upload has NOT been performed."
    )

    print(
        "\nThe approved state is now ready "
        "for the future Naukri upload module."
    )




def analyze_current_resume():

    resume = read_resume()

    baseline = load_baseline()

    current_fingerprint = calculate_fingerprint(
        resume["full_text"]
    )

    baseline_fingerprint = baseline[
        "fingerprint"
    ]

    print("\nCurrent fingerprint:")
    print(current_fingerprint)

    print("\nBaseline fingerprint:")
    print(baseline_fingerprint)

    # --------------------------------------------------------
    # No changes
    # --------------------------------------------------------

    if current_fingerprint == baseline_fingerprint:

        print("\n========================================")
        print("NO RESUME CHANGES DETECTED")
        print("========================================")

        print(
            "\nNo Naukri update is required."
        )

        analysis = {
            "changed": False,
            "meaningful_change_count": 0,
            "format_only_change_count": 0,
            "review_required_count": 0,
            "changes": [],
            "approval_status": "NOT_REQUIRED",
            "naukri_upload_approved": False,
        }

        ANALYSIS_FILE.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        with ANALYSIS_FILE.open(
                "w",
                encoding="utf-8",
        ) as f:

            json.dump(
                analysis,
                f,
                indent=2,
                ensure_ascii=False,
            )

        return

    # --------------------------------------------------------
    # Changes detected
    # --------------------------------------------------------

    changes = detect_changes(
        baseline["full_text"],
        resume["full_text"],
    )

    analysis = analyze_changes(
        changes
    )

    result = save_analysis(
        current_fingerprint,
        baseline_fingerprint,
        changes,
        analysis,
    )

    print("\n========================================")
    print("RESUME CHANGE ANALYSIS")
    print("========================================")

    for item in analysis:

        print(
            f"\nLine {item['line']}"
        )

        print(
            f"Classification: "
            f"{item['classification']}"
        )

        print(
            f"Reason: "
            f"{item['reason']}"
        )

        if item["old"]:
            print(
                f"OLD: {item['old']}"
            )
        else:
            print(
                "OLD: <new line>"
            )

        if item["new"]:
            print(
                f"NEW: {item['new']}"
            )
        else:
            print(
                "NEW: <removed line>"
            )

    print("\n========================================")
    print("APPROVAL STATUS")
    print("========================================")

    print(
        f"Meaningful changes: "
        f"{result['meaningful_change_count']}"
    )

    print(
        f"Review required: "
        f"{result['review_required_count']}"
    )

    print(
        f"Format-only changes: "
        f"{result['format_only_count']}"
    )

    print(
        "\nApproval status: PENDING"
    )

    print(
        "\nNo Naukri upload performed."
    )

def verify_upload_approval():

    print("\n========================================")
    print("NAUKRI UPLOAD APPROVAL VERIFICATION")
    print("========================================")

    try:
        analysis = load_analysis()
    except FileNotFoundError as e:
        print(f"\nERROR: {e}")
        return False

    approval_status = analysis.get(
        "approval_status",
        "UNKNOWN",
    )

    upload_approved = analysis.get(
        "naukri_upload_approved",
        False,
    )

    approved_at = analysis.get(
        "approved_at",
        "N/A",
    )

    changed = analysis.get(
        "changed",
        False,
    )

    meaningful_count = analysis.get(
        "meaningful_change_count",
        0,
    )

    review_count = analysis.get(
        "review_required_count",
        0,
    )

    print(
        f"\nResume changed       : {changed}"
    )

    print(
        f"Meaningful changes   : {meaningful_count}"
    )

    print(
        f"Review required      : {review_count}"
    )

    print(
        f"Approval status      : {approval_status}"
    )

    print(
        f"Naukri upload approved: "
        f"{upload_approved}"
    )

    print(
        f"Approved at          : {approved_at}"
    )

    # --------------------------------------------------------
    # Safety gate
    # --------------------------------------------------------

    if (
            approval_status == "APPROVED"
            and upload_approved is True
    ):

        print("\n========================================")
        print("UPLOAD GATE: APPROVED")
        print("========================================")

        print(
            "\nThe resume is explicitly approved "
            "for the future Naukri upload module."
        )

        return True

    print("\n========================================")
    print("UPLOAD GATE: BLOCKED")
    print("========================================")

    print(
        "\nNaukri upload must NOT be performed."
    )

    return False


def main():

    print("========================================")
    print("Naukri Resume Manager - Step 4")
    print("Approval Workflow")
    print("========================================")

    # --------------------------------------------------------
    # Approval command
    # --------------------------------------------------------

    if (
            len(sys.argv) > 1
            and sys.argv[1].lower()
            == "approve"
    ):

        approve_current_change()
        return
    # --------------------------------------------------------
    # Verification command
    # --------------------------------------------------------
    if (
            len(sys.argv) > 1
            and sys.argv[1].lower()
            == "verify"
    ):

        verify_upload_approval()
        return

    # --------------------------------------------------------
    # Normal analysis
    # --------------------------------------------------------

    analyze_current_resume()

    print("\n========================================")
    print("STEP 4 COMPLETE")
    print("========================================")

if __name__ == "__main__":
    main()
